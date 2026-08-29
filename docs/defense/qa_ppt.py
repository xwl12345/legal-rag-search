# -*- coding: utf-8 -*-
"""PPT QA：文本溢出估算 / 越界 / 文本框重叠检查"""
import math, sys
from pptx import Presentation
from pptx.util import Emu

PATH = sys.argv[1] if len(sys.argv) > 1 else "docs/defense/法律RAG检索引擎-答辩PPT.pptx"
EMU_IN = 914400.0
SLIDE_W, SLIDE_H = 13.333, 7.5

def char_units(s):
    u = 0.0
    for ch in s:
        u += 1.0 if ord(ch) > 0x2E80 else 0.55
    return u

def para_font(p, default=18.0):
    sz = None
    for r in p.runs:
        if r.font.size:
            sz = max(sz or 0, r.font.size.pt)
    return sz or default

issues = []
prs = Presentation(PATH)
for si, slide in enumerate(prs.slides, 1):
    texts = []  # (name, l, t, w, h) 只统计有字的文本框
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        tf = sh.text_frame
        txt = tf.text.strip()
        l, t = sh.left / EMU_IN, sh.top / EMU_IN
        w, h = sh.width / EMU_IN, sh.height / EMU_IN
        name = f"S{si} '{(txt[:14] + '..') if len(txt) > 14 else txt}' @({l:.2f},{t:.2f} {w:.2f}x{h:.2f})"
        if not txt:
            continue
        texts.append((name, l, t, w, h))
        # 越界
        if l < -0.02 or t < -0.02 or l + w > SLIDE_W + 0.02 or t + h > SLIDE_H + 0.02:
            issues.append(f"[越界] {name}")
        # 溢出估算
        w_pt, h_pt = w * 72, h * 72
        need = 0.0
        for p in tf.paragraphs:
            fs = para_font(p)
            cpu = max(1.0, w_pt / (fs * 1.06))  # 每行字符单位容量
            lines = 0
            for seg in p.text.split("\n"):
                lines += max(1, math.ceil(char_units(seg) / cpu)) if seg else 1
            need += lines * fs * 1.30
        need_in = need / 72.0
        if need_in > h * 1.10 + 0.03:
            issues.append(f"[溢出] {name} 需≈{need_in:.2f}in > 框{h:.2f}in")
    # 文本框重叠（同页两两）
    for i in range(len(texts)):
        for j in range(i + 1, len(texts)):
            a, b = texts[i], texts[j]
            ox = max(0, min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1]))
            oy = max(0, min(a[2] + a[4], b[2] + b[4]) - max(a[2], b[2]))
            if ox > 0.05 and oy > 0.05:
                issues.append(f"[重叠] {a[0]} × {b[0]} ({ox:.2f}x{oy:.2f}in)")

print(f"共 {len(prs.slides)} 页")
if issues:
    print(f"发现 {len(issues)} 个问题：")
    for i in issues:
        print(" ", i)
else:
    print("QA 通过：无溢出/越界/重叠")
